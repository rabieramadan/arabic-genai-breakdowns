#!/usr/bin/env python3
"""Apply journal-style formatting to the pandoc-generated DOCX.

Times New Roman throughout (including complex-script and East Asian slots, so Arabic
transliteration and Greek symbols inherit it too), bold title and headings, justified body
text, centred figures, and a gridded table style.

  python format_docx.py Paper2_revised.docx [--size 12] [--line 2.0]
"""
import argparse
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "Times New Roman"
HEADING_STYLES = {"Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3",
                  "Heading 4", "Heading 5", "Heading 6"}
CENTRED = {"Image Caption", "Figure"}


def set_font(run_or_style, size=None, bold=None):
    f = run_or_style.font
    f.name = FONT
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    rpr = (run_or_style._element.get_or_add_rPr() if hasattr(run_or_style._element, "get_or_add_rPr")
           else run_or_style._element.rPr)
    if rpr is not None:
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr.makeelement(qn("w:rFonts"), {}); rpr.append(rf)
        for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(slot), FONT)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--size", type=float, default=12.0, help="body point size")
    ap.add_argument("--line", type=float, default=None, help="line spacing, e.g. 2.0 for double")
    a = ap.parse_args()

    doc = Document(a.path)

    # document-wide default, so any style that inherits picks up the face
    set_font(doc.styles["Normal"], size=a.size)
    for st in doc.styles:
        try:
            if st.type is not None and st.font is not None:
                set_font(st)
                if st.name in HEADING_STYLES:
                    st.font.bold = True
                    st.font.color.rgb = RGBColor(0, 0, 0)
        except (AttributeError, NotImplementedError):
            continue

    body = heads = figs = 0
    for p in doc.paragraphs:
        name = p.style.name
        for r in p.runs:
            set_font(r)
        has_image = bool(p._element.findall(".//" + qn("w:drawing")))
        if name in HEADING_STYLES:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.bold = True
            heads += 1
        elif has_image or name in CENTRED:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            figs += 1
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if a.line:
                p.paragraph_format.line_spacing = a.line
            body += 1

    for t in doc.tables:
        try:
            t.style = doc.styles["Table Grid"]
        except KeyError:
            # pandoc's output has no Table Grid style; write the borders directly
            tbl_pr = t._tbl.tblPr
            borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                el = borders.makeelement(qn("w:" + edge), {})
                el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4")
                el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
                borders.append(el)
            tbl_pr.append(borders)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:
                        set_font(r, size=max(9.0, a.size - 2))
        for cell in t.rows[0].cells:          # header row bold
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True

    doc.save(a.path)
    print(f"{a.path}: {heads} headings bold, {body} paragraphs justified, "
          f"{figs} figure/caption paragraphs centred, {len(doc.tables)} tables gridded, font {FONT} {a.size}pt")


if __name__ == "__main__":
    main()
